#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
utils/report_generator.py - 스캔 결과 리포트 생성 (JSON, HTML, Excel)
"""

import json
import os
from datetime import datetime
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    print("[!] openpyxl이 설치되지 않아 Excel 리포트를 생성할 수 없습니다.")
    print("[!] 설치: pip install openpyxl")


# 중요도 매핑
SEVERITY_MAP = {
    '1. 버퍼 오버플로우': '상',
    '2. 포맷스트링': '상',
    '3. LDAP 인젝션': '상',
    '4. 운영체제 명령 실행': '상',
    '5. SQL 인젝션': '상',
    '6. SSI 인젝션': '상',
    '7. XPath 인젝션': '상',
    '8. 디렉터리 인덱싱': '상',
    '9. 정보 누출': '상',
    '10. 악성 콘텐츠': '상',
    '11. XSS (크로스사이트 스크립팅)': '상',
    '12. 약한 문자열 강도': '상',
    '13. 불충분한 인증': '상',
    '14. 취약한 패스워드 복구': '상',
    '15. CSRF': '상',
    '16. 세션 예측': '상',
    '17. 불충분한 인가': '상',
    '18. 불충분한 세션 만료': '상',
    '19. 세션 고정': '상',
    '20. 자동화 공격': '상',
    '21. 프로세스 검증 누락': '상',
    '22. 파일 업로드': '상',
    '23. 파일 다운로드': '상',
    '24. 관리자 페이지 노출': '상',
    '26. 위치 공개': '상',
    '27. 데이터 평문 전송': '상',
    '28. 쿠키 변조': '상',
}


def print_banner(target_url):
    """스캔 시작 배너 출력"""
    print(f"\n{'='*70}")
    print(f"웹 취약점 진단 시작")
    print(f"{'='*70}")
    print(f"대상 URL: {target_url}")
    print(f"시작 시간: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"{'='*70}")


def generate_report(target_url, results):
    """JSON 리포트 생성 및 저장"""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    # 요약 통계
    status_count = {'취약': 0, '양호': 0, '인터뷰': 0}
    for result in results.values():
        status_count[result['status']] += 1
    
    # 콘솔 출력
    print(f"\n{'='*70}")
    print(f"웹 취약점 진단 결과 요약")
    print(f"{'='*70}")
    
    print(f"\n📊 통계")
    print(f"   🔴 취약: {status_count['취약']}개")
    print(f"   🟢 양호: {status_count['양호']}개")
    print(f"   🟡 인터뷰: {status_count['인터뷰']}개")
    
    print(f"\n📋 상세 결과")
    for test_name, result in results.items():
        status = result['status']
        color = '🔴' if status == '취약' else '🟢' if status == '양호' else '🟡'
        print(f"\n{color} [{status}] {test_name}")
        print(f"   {result['description']}")
        
        if result.get('details'):
            for detail in result['details'][:3]:
                print(f"   - {detail}")
            if len(result['details']) > 3:
                print(f"   ... 외 {len(result['details']) - 3}개")
    
    print(f"\n{'='*70}")
    
    # JSON 리포트 생성
    report_data = {
        'target_url': target_url,
        'scan_time': timestamp,
        'summary': status_count,
        'results': results
    }
    
    # reports 디렉터리 생성
    os.makedirs('reports', exist_ok=True)
    
    # 파일명 생성
    filename = f"reports/scan_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    
    # JSON 파일 저장
    with open(filename, 'w', encoding='utf-8') as f:
        json.dump(report_data, f, ensure_ascii=False, indent=2)
    
    return filename


def generate_html_report(target_url, results):
    """HTML 형식의 리포트 생성"""
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    # 통계 계산
    status_count = {'취약': 0, '양호': 0, '인터뷰': 0}
    for result in results.values():
        status_count[result['status']] += 1
    
    html_template = f"""<!DOCTYPE html>
<html lang="ko">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>웹 취약점 진단 리포트</title>
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{ 
            font-family: 'Malgun Gothic', 'Apple SD Gothic Neo', sans-serif; 
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            padding: 20px;
            line-height: 1.6;
        }}
        .container {{ 
            max-width: 1200px; 
            margin: 0 auto; 
            background: white; 
            padding: 40px; 
            box-shadow: 0 10px 40px rgba(0,0,0,0.2);
            border-radius: 10px;
        }}
        h1 {{ 
            color: #2d3748; 
            border-bottom: 4px solid #667eea; 
            padding-bottom: 15px;
            margin-bottom: 30px;
            font-size: 2em;
        }}
        .info {{ 
            background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
            color: white;
            padding: 20px; 
            border-radius: 8px; 
            margin: 20px 0;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
        }}
        .info strong {{ font-size: 1.1em; }}
        .summary {{
            display: grid;
            grid-template-columns: repeat(3, 1fr);
            gap: 20px;
            margin: 30px 0;
        }}
        .summary-box {{
            padding: 25px;
            border-radius: 8px;
            text-align: center;
            box-shadow: 0 4px 6px rgba(0,0,0,0.1);
            transition: transform 0.3s;
        }}
        .summary-box:hover {{ transform: translateY(-5px); }}
        .summary-box h3 {{ font-size: 2.5em; margin: 10px 0; }}
        .summary-box p {{ font-size: 1.2em; color: #4a5568; }}
        .vulnerable {{ background: #fed7d7; border-left: 6px solid #f56565; }}
        .safe {{ background: #c6f6d5; border-left: 6px solid #48bb78; }}
        .interview {{ background: #feebc8; border-left: 6px solid #ed8936; }}
        .result-item {{
            padding: 20px;
            margin: 15px 0;
            border-radius: 8px;
            box-shadow: 0 2px 4px rgba(0,0,0,0.1);
            transition: all 0.3s;
        }}
        .result-item:hover {{ box-shadow: 0 4px 8px rgba(0,0,0,0.15); }}
        .status {{ 
            font-weight: bold; 
            font-size: 1.3em; 
            margin-bottom: 10px;
            display: flex;
            align-items: center;
            gap: 10px;
        }}
        .badge {{
            display: inline-block;
            padding: 5px 12px;
            border-radius: 20px;
            font-size: 0.8em;
            font-weight: bold;
            color: white;
        }}
        .badge-high {{ background: #f56565; }}
        .badge-medium {{ background: #ed8936; }}
        .badge-low {{ background: #48bb78; }}
        .details {{ 
            margin: 15px 0 0 30px; 
            color: #4a5568;
            background: #f7fafc;
            padding: 15px;
            border-radius: 5px;
        }}
        .details ul {{ list-style-type: none; padding: 0; }}
        .details li {{ 
            padding: 8px 0; 
            border-bottom: 1px solid #e2e8f0;
            position: relative;
            padding-left: 20px;
        }}
        .details li:before {{
            content: "▸";
            position: absolute;
            left: 0;
            color: #667eea;
            font-weight: bold;
        }}
        .details li:last-child {{ border-bottom: none; }}
        .footer {{
            margin-top: 40px;
            padding-top: 20px;
            border-top: 2px solid #e2e8f0;
            text-align: center;
            color: #718096;
        }}
    </style>
</head>
<body>
    <div class="container">
        <h1>🔍 웹 취약점 진단 리포트</h1>
        
        <div class="info">
            <strong>대상 URL:</strong> {target_url}<br>
            <strong>스캔 시간:</strong> {timestamp}
        </div>
        
        <h2 style="margin: 30px 0 20px 0; color: #2d3748;">📊 검사 결과 요약</h2>
        <div class="summary">
            <div class="summary-box" style="background: #fed7d7;">
                <p>취약</p>
                <h3 style="color: #c53030;">{status_count['취약']}</h3>
            </div>
            <div class="summary-box" style="background: #c6f6d5;">
                <p>양호</p>
                <h3 style="color: #2f855a;">{status_count['양호']}</h3>
            </div>
            <div class="summary-box" style="background: #feebc8;">
                <p>인터뷰</p>
                <h3 style="color: #c05621;">{status_count['인터뷰']}</h3>
            </div>
        </div>
        
        <h2 style="margin: 30px 0 20px 0; color: #2d3748;">📋 상세 검사 결과</h2>
"""
    
    for test_name, result in results.items():
        status = result['status']
        css_class = 'vulnerable' if status == '취약' else 'safe' if status == '양호' else 'interview'
        emoji = '🔴' if status == '취약' else '🟢' if status == '양호' else '🟡'
        
        # 중요도 뱃지
        severity = SEVERITY_MAP.get(test_name, '중')
        badge_class = 'badge-high' if severity == '상' else 'badge-medium' if severity == '중' else 'badge-low'
        severity_text = '높음' if severity == '상' else '보통' if severity == '중' else '낮음'
        
        html_template += f"""
        <div class="result-item {css_class}">
            <div class="status">
                <span>{emoji}</span>
                <span>[{status}] {test_name}</span>
                <span class="badge {badge_class}">중요도: {severity_text}</span>
            </div>
            <p style="margin: 10px 0; color: #2d3748; font-size: 1.1em;">{result['description']}</p>
"""
        
        if result.get('details'):
            html_template += "<div class='details'><ul>"
            for detail in result['details']:
                html_template += f"<li>{detail}</li>"
            html_template += "</ul></div>"
        
        html_template += "</div>"
    
    html_template += """
        <div class="footer">
            <p><strong>웹 취약점 자동 진단 스크립트 v2.0</strong></p>
            <p>이 리포트는 자동화 도구로 생성되었으며, 일부 결과는 수동 확인이 필요할 수 있습니다.</p>
        </div>
    </div>
</body>
</html>
"""
    
    # HTML 파일 저장
    os.makedirs('reports', exist_ok=True)
    filename = f"reports/scan_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
    
    with open(filename, 'w', encoding='utf-8') as f:
        f.write(html_template)
    
    return filename


def generate_excel_report(target_url, results):
    """Excel 형식의 리포트 생성"""
    
    if not EXCEL_AVAILABLE:
        print("[!] Excel 리포트를 생성할 수 없습니다. openpyxl을 설치하세요.")
        return None
    
    try:
        # 워크북 생성
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "취약점 진단 결과"
        
        # 스타일 정의
        header_font = Font(name='맑은 고딕', size=12, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        cell_alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        center_alignment = Alignment(horizontal='center', vertical='center')
        
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # 제목 행
        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = f'웹 취약점 진단 리포트 - {target_url}'
        title_cell.font = Font(name='맑은 고딕', size=14, bold=True)
        title_cell.alignment = center_alignment
        title_cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
        
        # 스캔 정보
        ws.merge_cells('A2:E2')
        info_cell = ws['A2']
        info_cell.value = f'스캔 시간: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}'
        info_cell.font = Font(name='맑은 고딕', size=10)
        info_cell.alignment = center_alignment
        
        # 빈 행
        ws.append([])
        
        # 헤더 행
        headers = ['코드', '항목명', '중요도', '결과', '결과 현황']
        ws.append(headers)
        
        # 헤더 스타일 적용
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # 열 너비 설정
        ws.column_dimensions['A'].width = 8   # 코드
        ws.column_dimensions['B'].width = 30  # 항목명
        ws.column_dimensions['C'].width = 10  # 중요도
        ws.column_dimensions['D'].width = 12  # 결과
        ws.column_dimensions['E'].width = 80  # 결과 현황
        
        # 데이터 행 추가
        row_num = 5
        for test_name, result in results.items():
            # 코드 추출 (예: "1. SQL 인젝션" -> "1")
            code = test_name.split('.')[0].strip()
            
            # 항목명 추출
            item_name = '.'.join(test_name.split('.')[1:]).strip()
            
            # 중요도
            severity = SEVERITY_MAP.get(test_name, '중')
            
            # 결과
            status = result['status']
            
            # 결과 현황 생성
            status_detail = result['description']
            if result.get('details'):
                status_detail += '\n\n상세 내용:\n'
                for i, detail in enumerate(result['details'], 1):
                    status_detail += f"{i}. {detail}\n"
            
            # 행 추가
            ws.append([code, item_name, severity, status, status_detail])
            
            # 스타일 적용
            for col_num in range(1, 6):
                cell = ws.cell(row=row_num, column=col_num)
                cell.border = border
                
                if col_num == 5:  # 결과 현황
                    cell.alignment = cell_alignment
                else:
                    cell.alignment = center_alignment if col_num in [1, 3, 4] else Alignment(horizontal='left', vertical='center')
                
                # 결과에 따른 색상
                if col_num == 4:  # 결과 열
                    if status == '취약':
                        cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='9C0006')
                    elif status == '양호':
                        cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='006100')
                    else:  # 인터뷰
                        cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='9C6500')
                else:
                    cell.font = Font(name='맑은 고딕', size=10)
            
            # 행 높이 자동 조정 (결과 현황이 긴 경우)
            ws.row_dimensions[row_num].height = max(30, len(status_detail.split('\n')) * 15)
            
            row_num += 1
        
        # 요약 통계 추가
        row_num += 2
        ws.merge_cells(f'A{row_num}:E{row_num}')
        summary_cell = ws.cell(row=row_num, column=1)
        
        status_count = {'취약': 0, '양호': 0, '인터뷰': 0}
        for result in results.values():
            status_count[result['status']] += 1
        
        summary_cell.value = f"총 {len(results)}개 항목 검사 완료 | 취약: {status_count['취약']}개 | 양호: {status_count['양호']}개 | 인터뷰: {status_count['인터뷰']}개"
        summary_cell.font = Font(name='맑은 고딕', size=11, bold=True)
        summary_cell.alignment = center_alignment
        summary_cell.fill = PatternFill(start_color='E7E6E6', end_color='E7E6E6', fill_type='solid')
        summary_cell.border = border
        
        # 파일 저장
        os.makedirs('reports', exist_ok=True)
        filename = f"reports/scan_report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        
        return filename
    
    except Exception as e:
        print(f"[!] Excel 리포트 생성 중 오류: {str(e)}")
        return None